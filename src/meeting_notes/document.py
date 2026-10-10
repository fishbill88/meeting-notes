"""Document generation module – saves meeting notes as Markdown or DOCX files."""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
from typing import Optional

from meeting_notes.summarizer import MeetingNotes


def generate_document(
    notes: MeetingNotes,
    transcript: Optional[str] = None,
    output_dir: str = ".",
    output_format: str = "markdown",
    filename: Optional[str] = None,
) -> str:
    """Write meeting notes to a file.

    Args:
        notes: Structured :class:`~meeting_notes.summarizer.MeetingNotes`.
        transcript: Optional full transcript text to append to the document.
        output_dir: Directory where the file will be saved.
        output_format: ``"markdown"`` (produces a ``.md`` file) or ``"docx"``
            (produces a ``.docx`` Word document).
        filename: Optional base filename (without extension).  A timestamped
            name is generated when not provided.

    Returns:
        Absolute path to the generated file.
    """
    os.makedirs(output_dir, exist_ok=True)

    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if filename is None:
        safe_title = _safe_filename(notes.title)
        filename = f"{timestamp}_{safe_title}"

    if output_format == "docx":
        out_path = os.path.join(output_dir, f"{filename}.docx")
        _write_docx(notes, transcript, out_path, date_str)
    else:
        out_path = os.path.join(output_dir, f"{filename}.md")
        _write_markdown(notes, transcript, out_path, date_str)

    return str(Path(out_path).resolve())


# ------------------------------------------------------------------
# Markdown writer
# ------------------------------------------------------------------


def _write_markdown(
    notes: MeetingNotes,
    transcript: Optional[str],
    path: str,
    date_str: str,
) -> None:
    content = notes.to_markdown(transcript=transcript, date_str=date_str)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


# ------------------------------------------------------------------
# DOCX writer
# ------------------------------------------------------------------


def _write_docx(
    notes: MeetingNotes,
    transcript: Optional[str],
    path: str,
    date_str: str,
) -> None:
    _build_docx(notes, transcript, date_str).save(path)


def _build_docx(
    notes: MeetingNotes,
    transcript: Optional[str],
    date_str: str,
):
    """Build the Word document in memory.

    Separate from :func:`_write_docx` so the desktop app can export bytes
    straight to a native save dialog without a temp file on the way.
    """
    try:
        from docx import Document  # noqa: PLC0415
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required to generate DOCX files. "
            "Install it with: pip install python-docx"
        ) from exc

    doc = Document()

    # Title
    title_para = doc.add_heading(notes.title, level=0)

    # Metadata
    meta = doc.add_paragraph()
    meta.add_run("Date: ").bold = True
    meta.add_run(date_str)

    if notes.language and notes.language != "en":
        lang_para = doc.add_paragraph()
        lang_para.add_run("Language detected: ").bold = True
        lang_para.add_run(notes.language)

    doc.add_paragraph()  # spacer

    # Summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(notes.summary)

    # Key Topics
    if notes.key_topics:
        doc.add_heading("Key Topics", level=1)
        for topic in notes.key_topics:
            doc.add_paragraph(topic, style="List Bullet")

    # Decisions
    if notes.decisions:
        doc.add_heading("Decisions Made", level=1)
        for decision in notes.decisions:
            doc.add_paragraph(decision, style="List Bullet")

    # Action Items
    if notes.action_items:
        doc.add_heading("Action Items", level=1)
        for item in notes.action_items:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("☐ ").bold = True
            p.add_run(item)

    # Full Transcript
    if transcript:
        doc.add_heading("Full Transcript", level=1)
        # Split into paragraphs to avoid a single giant paragraph
        for chunk in transcript.split("\n"):
            if chunk.strip():
                doc.add_paragraph(chunk.strip())

    return doc


# ------------------------------------------------------------------
# In-memory rendering (desktop export)
# ------------------------------------------------------------------


#: Formats :func:`render_document` can produce, mapped to their MIME types.
EXPORT_FORMATS = {
    "markdown": "text/markdown",
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    "text": "text/plain",
}


def render_document(
    notes: MeetingNotes,
    transcript: Optional[str] = None,
    output_format: str = "markdown",
    date_str: Optional[str] = None,
) -> bytes:
    """Render notes to bytes without touching the filesystem.

    The desktop app hands these straight to a native save dialog, so the
    document never lands in a temp directory — which matters when the content
    is a patient consultation.

    Args:
        notes: Structured notes.
        transcript: Optional transcript to append.
        output_format: One of :data:`EXPORT_FORMATS`.
        date_str: Display date; defaults to now.

    Returns:
        Encoded document bytes.
    """
    if output_format not in EXPORT_FORMATS:
        raise ValueError(
            f"Invalid output_format '{output_format}'. "
            f"Choose from: {', '.join(EXPORT_FORMATS)}"
        )

    date_str = date_str or datetime.now().strftime("%Y-%m-%d %H:%M")

    if output_format == "docx":
        import io  # noqa: PLC0415

        buffer = io.BytesIO()
        _build_docx(notes, transcript, date_str).save(buffer)
        return buffer.getvalue()

    markdown = notes.to_markdown(transcript=transcript, date_str=date_str)
    if output_format == "text":
        return _strip_markdown(markdown).encode("utf-8")
    return markdown.encode("utf-8")


def suggested_filename(notes: MeetingNotes, output_format: str) -> str:
    """A timestamped, filesystem-safe filename for an export."""
    extension = {"markdown": "md", "docx": "docx", "text": "txt"}.get(
        output_format, "md"
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{timestamp}_{_safe_filename(notes.title)}.{extension}"


def _strip_markdown(text: str) -> str:
    """Flatten Markdown for pasting into an EMR field that renders nothing."""
    import re  # noqa: PLC0415

    out = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    out = re.sub(r"\*\*(.+?)\*\*", r"\1", out)
    out = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", out)
    out = re.sub(r"^\s*[-*]\s+", "• ", out, flags=re.MULTILINE)
    # Action items are written as task-list checkboxes. Rendered nowhere, they
    # read as literal "[ ]" noise in a pasted chart entry.
    out = re.sub(r"^(•\s*)\[[ xX]\]\s*", r"\1", out, flags=re.MULTILINE)
    return out


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------


def _safe_filename(text: str, max_len: int = 40) -> str:
    """Convert arbitrary text to a filesystem-safe filename fragment."""
    import re
    safe = re.sub(r"[^\w\s-]", "", text)
    safe = re.sub(r"[\s-]+", "_", safe).strip("_")
    return safe[:max_len] if safe else "meeting"
